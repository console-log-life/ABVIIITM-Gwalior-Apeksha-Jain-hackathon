"""Import the reviewed spreadsheet back into data/eval/labelled_headlines.csv, then re-run scripts/evaluate.py.

For every row: final label = the 'corrected' value if filled in, else the draft. Every value is validated
(sentiment, the 11 event classes, ticker format, label_status). Nothing is written if any row is invalid.
Rows are matched by id; the text and metadata in the CSV are kept (a text edited in the sheet is reported, not used).

Usage:  python scripts/import_label_review.py [--xlsx data/eval/label_review.xlsx]
                                             [--csv data/eval/labelled_headlines.csv] [--no-evaluate]
"""

from __future__ import annotations

import argparse
import csv
import os
import re
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from risk_engine.schemas import EVENT_TYPES  # noqa: E402

XLSX_DEFAULT = ROOT / "data" / "eval" / "label_review.xlsx"
CSV_DEFAULT = ROOT / "data" / "eval" / "labelled_headlines.csv"
SENTIMENTS = {"Negative", "Neutral", "Positive"}
STATUSES = {"draft_agent", "human_reviewed"}
TICKER_RE = re.compile(r"^(?:MARKET|UNRESOLVED|SOV-[A-Z]{2}|[A-Z0-9][A-Z0-9.\-]{0,14})$")


class ReviewError(ValueError):
    pass


def _clean(v) -> str:
    return "" if v is None else str(v).strip()


def import_review(xlsx: Path, csv_path: Path) -> dict[str, int]:
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb["Review"]
    it = ws.iter_rows(values_only=True)
    header = [_clean(h) for h in next(it)]
    idx = {h: i for i, h in enumerate(header)}
    need = ["id", "text", "draft_sentiment", "draft_event", "draft_ticker", "sentiment_corrected", "event_corrected",
            "ticker_corrected", "label_status", "reviewer_notes"]
    missing = [h for h in need if h not in idx]
    if missing:
        raise ReviewError(f"sheet 'Review' is missing columns {missing}")

    rows = list(csv.DictReader(open(csv_path, encoding="utf-8")))
    fields = list(rows[0].keys())
    by_id = {r["id"]: r for r in rows}
    errors, changed, reviewed, text_edits, seen = [], 0, 0, 0, set()
    for n, raw in enumerate(it, start=2):
        if raw is None or all(v is None for v in raw):
            continue
        rid = _clean(raw[idx["id"]])
        if rid.endswith(".0"):
            rid = rid[:-2]
        if rid not in by_id:
            errors.append(f"row {n}: unknown id {rid!r}")
            continue
        seen.add(rid)
        rec = by_id[rid]
        sent = _clean(raw[idx["sentiment_corrected"]]) or _clean(raw[idx["draft_sentiment"]])
        event = _clean(raw[idx["event_corrected"]]) or _clean(raw[idx["draft_event"]])
        ticker = (_clean(raw[idx["ticker_corrected"]]) or _clean(raw[idx["draft_ticker"]])).upper()
        status = _clean(raw[idx["label_status"]]) or "draft_agent"
        if sent not in SENTIMENTS:
            errors.append(f"row {n} (id {rid}): sentiment {sent!r} not in {sorted(SENTIMENTS)}")
        if event not in EVENT_TYPES:
            errors.append(f"row {n} (id {rid}): event {event!r} is not one of the 11 classes")
        if not TICKER_RE.match(ticker):
            errors.append(f"row {n} (id {rid}): ticker {ticker!r} must be a ticker, SOV-xx, MARKET or UNRESOLVED")
        if status not in STATUSES:
            errors.append(f"row {n} (id {rid}): label_status {status!r} not in {sorted(STATUSES)}")
        if _clean(raw[idx["text"]]) != rec["text"].strip():
            text_edits += 1
        new = {"gold_sentiment": sent, "gold_event": event, "gold_ticker": ticker, "label_status": status}
        if any(rec[k] != v for k, v in new.items() if k != "label_status"):
            changed += 1
        reviewed += status == "human_reviewed"
        notes = _clean(raw[idx["reviewer_notes"]])
        rec.update(new)
        if notes:
            rec["notes"] = f"reviewer: {notes}" + (f" | agent: {rec['notes']}" if rec["notes"] and
                                                   not rec["notes"].startswith("reviewer:") else "")
    absent = sorted(set(by_id) - seen, key=int)
    if absent:
        errors.append(f"{len(absent)} ids from the CSV are missing in the sheet: {absent[:10]}")
    if errors:
        raise ReviewError("nothing written; fix these rows:\n  " + "\n  ".join(errors[:40]))

    tmp = csv_path.with_suffix(".csv.tmp")
    with open(tmp, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)
    os.replace(tmp, csv_path)
    return {"rows": len(rows), "labels_changed": changed, "human_reviewed": reviewed, "text_edits_ignored": text_edits}


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", type=Path, default=XLSX_DEFAULT)
    ap.add_argument("--csv", type=Path, default=CSV_DEFAULT)
    ap.add_argument("--no-evaluate", action="store_true")
    args = ap.parse_args()
    try:
        stats = import_review(args.xlsx, args.csv)
    except ReviewError as exc:
        print(f"IMPORT FAILED: {exc}")
        return 1
    print(f"imported {stats['rows']} rows: {stats['labels_changed']} label rows changed, "
          f"{stats['human_reviewed']} marked human_reviewed"
          + (f", {stats['text_edits_ignored']} edited texts ignored" if stats["text_edits_ignored"] else ""))
    if args.no_evaluate:
        return 0
    label = f"after human review ({stats['human_reviewed']}/{stats['rows']} rows reviewed)"
    cmd = [sys.executable, str(ROOT / "scripts" / "evaluate.py"), "--label", label]
    return subprocess.run(cmd, cwd=ROOT).returncode


if __name__ == "__main__":
    sys.exit(main())
