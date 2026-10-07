"""Build data/eval/label_review.xlsx from data/eval/labelled_headlines.csv for fast HUMAN label review.

- One row per item: text, source, the DRAFT gold labels (written by an AI assistant), and empty correction columns.
- Dropdowns (data validation) for corrected sentiment, corrected event (the 11 classes exactly) and label_status.
- Model predictions are deliberately NOT included, so they cannot anchor the reviewer.
- A separate "How to review" sheet holds the legend and an example row (never imported).
Read it back with src/scripts/import_label_review.py.

Usage:  python src/scripts/build_label_review.py [--csv data/eval/labelled_headlines.csv]
                                            [--out data/eval/label_review.xlsx]
"""

from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src"))

from risk_engine.schemas import EVENT_TYPES  # noqa: E402

CSV_DEFAULT = ROOT / "data" / "eval" / "labelled_headlines.csv"
OUT_DEFAULT = ROOT / "data" / "eval" / "label_review.xlsx"
SENTIMENTS = ["Negative", "Neutral", "Positive"]
STATUSES = ["draft_agent", "human_reviewed"]
SHEET = "Review"
# (header, width, kind) kind: ro = read-only info, draft = AI draft, input = reviewer fills in
COLUMNS = [
    ("id", 6, "ro"), ("text", 70, "ro"), ("source", 12, "ro"), ("source_type", 11, "ro"),
    ("draft_sentiment", 14, "draft"), ("draft_event", 16, "draft"), ("draft_ticker", 15, "draft"),
    ("sentiment_corrected", 16, "input"), ("event_corrected", 17, "input"), ("ticker_corrected", 16, "input"),
    ("label_status", 16, "input"), ("reviewer_notes", 34, "input"), ("agent_notes", 34, "ro"),
]
FONT = "Arial"
NAVY = "1F2A44"
INPUT_FILL = PatternFill("solid", fgColor="FFF6CC")
DRAFT_FILL = PatternFill("solid", fgColor="EEF2F7")
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
THIN = Side(style="thin", color="D0D4DA")


def build(csv_path: Path, out: Path) -> int:
    rows = list(csv.DictReader(open(csv_path, encoding="utf-8")))
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append([c[0] for c in COLUMNS])
    for r in rows:
        ws.append([int(r["id"]), r["text"], r["source"], r["source_type"], r["gold_sentiment"], r["gold_event"],
                   r["gold_ticker"], None, None, None, r["label_status"], None, r["notes"] or None])

    body = Font(name=FONT, size=10)
    for i, (_, width, kind) in enumerate(COLUMNS, start=1):
        letter = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[letter].width = width
        head = ws.cell(row=1, column=i)
        head.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        head.fill = HEAD_FILL
        head.alignment = Alignment(wrap_text=True, vertical="center")
        for row in range(2, len(rows) + 2):
            c = ws.cell(row=row, column=i)
            c.font = body
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = Border(bottom=THIN)
            if kind == "input":
                c.fill = INPUT_FILL
            elif kind == "draft":
                c.fill = DRAFT_FILL
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "C2"  # header row + id/text stay visible
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(COLUMNS)).column_letter}{len(rows) + 1}"

    last = len(rows) + 1
    col = {name: ws.cell(row=1, column=i).column_letter for i, (name, _, _) in enumerate(COLUMNS, start=1)}
    validations = [
        (SENTIMENTS, col["sentiment_corrected"], "Pick a sentiment, or leave blank to accept the draft."),
        (list(EVENT_TYPES), col["event_corrected"], "Pick one of the 11 event classes, or leave blank."),
        (STATUSES, col["label_status"], "Set to human_reviewed once you have checked this row."),
    ]
    for values, letter, prompt in validations:
        dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True,
                            showErrorMessage=True, errorTitle="Invalid value",
                            error="Choose a value from the list.", promptTitle="Input", prompt=prompt)
        ws.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}{last}")

    guide = wb.create_sheet("How to review")
    lines = [
        ("How to review the evaluation labels", True),
        ("", False),
        ("1. Work on the 'Review' sheet. Read the text, then check the three DRAFT labels (grey columns). They were "
         "drafted by an AI assistant.", False),
        ("2. If a draft is wrong, choose the right value in the yellow 'corrected' column. Leave it blank to accept "
         "the draft.", False),
        ("3. Set label_status to human_reviewed for every row you have checked (whether or not you changed "
         "it).", False),
        ("4. ticker_corrected: a universe ticker (e.g. AAPL, TATAMOTORS.NS), a sovereign issuer id (SOV-US, "
         "SOV-IN), MARKET (market-wide story) or UNRESOLVED (company not in our universe).", False),
        ("5. Model predictions are intentionally not shown, so they cannot anchor your judgement.", False),
        ("6. When done: powershell -ExecutionPolicy Bypass -File tasks.ps1 import-labels", False),
        ("   This validates every value, writes data/eval/labelled_headlines.csv and re-runs src/scripts/evaluate.py.",
         False),
        ("", False),
        ("Example of a corrected row (format only; this sheet is never imported):", True),
    ]
    for text, bold in lines:
        guide.append([text])
        guide.cell(row=guide.max_row, column=1).font = Font(name=FONT, size=11 if not bold else 13, bold=bold)
    guide.append([c[0] for c in COLUMNS])
    for i in range(1, len(COLUMNS) + 1):
        c = guide.cell(row=guide.max_row, column=i)
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
    guide.append([999, "Example Corp shares jump after regulator approves merger", "google_news", "news", "Neutral",
                  "Regulatory", "UNRESOLVED", "Positive", "M&A", None, "human_reviewed",
                  "merger approval is the main event", None])
    for i, (_, width, kind) in enumerate(COLUMNS, start=1):
        c = guide.cell(row=guide.max_row, column=i)
        c.font = Font(name=FONT, size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if kind == "input":
            c.fill = INPUT_FILL
        elif kind == "draft":
            c.fill = DRAFT_FILL
        guide.column_dimensions[c.column_letter].width = max(width, 12) if i > 1 else 110
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return len(rows)


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--csv", type=Path, default=CSV_DEFAULT)
    ap.add_argument("--out", type=Path, default=OUT_DEFAULT)
    args = ap.parse_args()
    n = build(args.csv, args.out)
    print(f"wrote {n} rows to {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
